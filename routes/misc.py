# Standard library imports
from datetime import datetime, timedelta
import io
from pathlib import Path
from zoneinfo import ZoneInfo

# Third-party imports
from flask import Blueprint, render_template, request, redirect, url_for, flash, Response, current_app
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

# Local imports
from extensions import db
from models import Product, ProductHistory, WarehouseProductBatch, StoreProductBatch, Requisition, RequisitionItem, ActivityLog
from helpers import user_redirect_to_requisition, admin_required, log_activity

# Create blueprint with URL prefix
misc = Blueprint('misc', __name__, url_prefix='/misc')


@misc.route('/')
@login_required
@user_redirect_to_requisition
def misc_dashboard():
    """Misc dashboard with transaction history and barcode download"""
    today = datetime.now(ZoneInfo("Asia/Manila")).date()
    # Calculate one month ago from today
    one_month_ago = today - timedelta(days=30)
    
    return render_template('shared/misc_dashboard.html', today=today, one_month_ago=one_month_ago)


@misc.route('/transaction_history')
@login_required
@user_redirect_to_requisition
def transaction_history():
    """Transaction history download page (legacy - redirects to dashboard)"""
    return redirect(url_for('misc.misc_dashboard'))


@misc.route('/download_transaction_history', methods=['POST'])
@login_required
@admin_required
def download_transaction_history():
    """Download transaction history as Excel file"""
    try:
        # Get date range from form
        from_date_str = request.form.get('from_date')
        to_date_str = request.form.get('to_date')
        
        if not from_date_str or not to_date_str:
            flash('Please select both from and to dates.', 'error')
            return redirect(url_for('misc.transaction_history'))
        
        # Parse dates
        from_date = datetime.strptime(from_date_str, '%Y-%m-%d').date()
        to_date = datetime.strptime(to_date_str, '%Y-%m-%d').date()
        
        if from_date > to_date:
            flash('From date cannot be later than to date.', 'error')
            return redirect(url_for('misc.transaction_history'))
        
        # Query transaction history within date range
        # Convert dates to datetime for proper comparison
        from_datetime = datetime.combine(from_date, datetime.min.time())
        to_datetime = datetime.combine(to_date, datetime.max.time())
        
        # Filter by company for Company-Admin users
        if current_user.role == 'Company-Admin' and current_user.company_id:
            # Company-Admin can only see their own company's transaction history
            history_records = ProductHistory.query.join(Product).filter(
                ProductHistory.trxn_date >= from_datetime,
                ProductHistory.trxn_date <= to_datetime,
                Product.company_id == current_user.company_id
            ).order_by(ProductHistory.product_id, ProductHistory.trxn_date).all()
        else:
            # SYS-Admin and WH-Admin can see all transaction history
            history_records = ProductHistory.query.filter(
                ProductHistory.trxn_date >= from_datetime,
                ProductHistory.trxn_date <= to_datetime
            ).order_by(ProductHistory.product_id, ProductHistory.trxn_date).all()
        
        if not history_records:
            flash('No transaction history found for the selected date range.', 'warning')
            return redirect(url_for('misc.transaction_history'))
        
        # Create Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Transaction History"
        
        # Define styles
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        title_font = Font(bold=True, size=14)
        title_alignment = Alignment(horizontal="center", vertical="center")
        
        # Define borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Write title
        ws.merge_cells('A1:G1')
        title_cell = ws.cell(row=1, column=1, value="ICEBERGS WAREHOUSE DETAILED INVENTORY")
        title_cell.font = title_font
        title_cell.alignment = title_alignment
        
        # Write date range
        ws.merge_cells('A2:G2')
        date_range = f"Period: {from_date.strftime('%b/%d/%Y')} to {to_date.strftime('%b/%d/%Y')}"
        date_cell = ws.cell(row=2, column=1, value=date_range)
        date_cell.font = Font(bold=True, size=12)
        date_cell.alignment = title_alignment
        
        # Write header row
        headers = [
            'item_description', 'trxn dt', 'trxn code', 'payee', 'IN', 'OUT', 'Total'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Group transactions by product
        current_row = 5
        current_product = None
        product_total = 0
        
        for record in history_records:
            # Get product details
            product = Product.query.get(record.product_id)
            if not product:
                continue
            
            # Format item description
            item_description = f"{product.name} ({product.uom or 'unit'})"
            
            # Format transaction date
            trxn_date = record.trxn_date.strftime('%m/%d/%Y')
            
            # Determine transaction code and payee
            if record.trxn_code == 'TRANS IN':
                trxn_code = 'TRANS IN'
                payee = record.payee or 'Icebergs Food Concepts, Inc.'
                in_qty = record.quantity_in
                out_qty = 0
                total = record.net_change
            elif record.trxn_code == 'TRANS OUT':
                trxn_code = 'TRANS OUT'
                payee = record.payee or 'ICEBERGS MARASBARAS'
                in_qty = 0
                out_qty = record.quantity_out
                total = record.net_change
            else:
                # Handle other transaction types
                trxn_code = record.trxn_code
                payee = record.payee or 'COOL ICE FOOD DESSERTS'
                in_qty = record.quantity_in if record.quantity_in else 0
                out_qty = record.quantity_out if record.quantity_out else 0
                total = record.net_change
            
            # Check if this is a new product (first transaction for this product)
            is_new_product = current_product != record.product_id
            
            # Write transaction row
            row_data = [
                item_description if is_new_product else '',  # Only show item description for first transaction
                trxn_date,
                trxn_code,
                payee,
                in_qty,
                out_qty,
                total
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col, value=value)
                cell.border = thin_border
                
                # Format numbers
                if col in [5, 6, 7]:  # IN, OUT, Total columns
                    if isinstance(value, (int, float)):
                        cell.number_format = '0.00'
            
            current_row += 1
            product_total += total
            current_product = record.product_id
            
            # Check if next record is for a different product
            next_record = history_records[history_records.index(record) + 1] if history_records.index(record) + 1 < len(history_records) else None
            
            if not next_record or next_record.product_id != record.product_id:
                # Write product total row
                total_row_data = [
                    f"{item_description} Total",
                    '',
                    '',
                    '',
                    '',
                    '',
                    product_total
                ]
                
                for col, value in enumerate(total_row_data, 1):
                    cell = ws.cell(row=current_row, column=col, value=value)
                    cell.border = thin_border
                    cell.font = Font(bold=True)
                    
                    if col == 7:  # Total column
                        cell.number_format = '0.00'
                
                current_row += 2  # Add extra space between products
                product_total = 0
        
        # Auto-adjust column widths
        column_widths = [40, 12, 15, 30, 10, 10, 12]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Prepare response
        timestamp = datetime.now(ZoneInfo("Asia/Manila")).strftime('%Y%m%d_%H%M%S')
        filename = f"transaction_history_{from_date.strftime('%Y%m%d')}_to_{to_date.strftime('%Y%m%d')}_{timestamp}.xlsx"
        
        # Log the export
        log_activity(
            action="Export Transaction History",
            details=f"Exported transaction history from {from_date.strftime('%Y-%m-%d')} to {to_date.strftime('%Y-%m-%d')} with {len(history_records)} records"
        )
        
        return Response(
            output.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment;filename={filename}"}
        )
        
    except Exception as e:
        # Log the error
        log_activity(
            action="Export Transaction History Error",
            details=f"Failed to export transaction history: {str(e)}"
        )
        flash(f'Error exporting transaction history: {str(e)}', 'error')
        return redirect(url_for('misc.transaction_history'))


def draw_dashed_rect(canvas, x, y, width, height, dash_length=2*mm, gap_length=1*mm):
    """Draw a dashed rectangle for easy cutting"""
    canvas.saveState()
    canvas.setDash(dash_length, gap_length)
    canvas.rect(x, y, width, height)
    canvas.restoreState()


def calculate_layout(page_width, page_height, barcode_width, barcode_height, margin=5*mm):
    """Calculate optimal layout for barcodes on a page"""
    available_width = page_width - 2 * margin
    available_height = page_height - 2 * margin
    
    # Calculate how many barcodes can fit horizontally and vertically
    cols = int(available_width // barcode_width)
    rows = int(available_height // barcode_height)
    
    # Ensure at least 1 column and row
    cols = max(1, cols)
    rows = max(1, rows)
    
    # Calculate spacing to center the grid
    spacing_x = (available_width - (cols * barcode_width)) / (cols + 1)
    spacing_y = (available_height - (rows * barcode_height)) / (rows + 1)
    
    return cols, rows, spacing_x, spacing_y


@misc.route('/download_barcodes')
@login_required
@admin_required
def download_barcodes():
    """Download barcode PDF formatted in a clean, non-overlapping grid layout without edge cropping."""
    try:
        from helpers import ensure_product_barcode
        
        # Get products from database, filtered by company for users with company assignment
        if hasattr(current_user, 'company_id') and current_user.company_id:
            products = Product.query.filter(
                Product.company_id == current_user.company_id
            ).order_by(Product.name).all()
        else:
            products = Product.query.order_by(Product.name).all()
        
        if not products:
            flash('No products found in database!', 'warning')
            return redirect(url_for('misc.transaction_history'))
        
        # Ensure all retrieved products have a barcode and PNG image generated
        for p in products:
            ensure_product_barcode(p)
            
        # Create PDF in memory
        output = io.BytesIO()
        
        # Standard A4: 210mm x 297mm
        page_width, page_height = A4
        margin_x = 8 * mm
        margin_top = 10 * mm
        margin_bottom = 12 * mm
        
        # Clean 3x7 Grid (21 labels per page - perfectly fits within printable boundaries)
        cols = 3
        rows = 7
        gap_x = 4 * mm
        gap_y = 3.5 * mm
        
        # Label dimensions: exact fit within available space
        avail_width = page_width - (2 * margin_x)
        label_width = (avail_width - (cols - 1) * gap_x) / cols  # 62.0 mm
        label_height = 36 * mm  # Generous height for name, category, and sharp barcode
        
        labels_per_page = cols * rows
        total_pages = (len(products) + labels_per_page - 1) // labels_per_page
        
        c = canvas.Canvas(output, pagesize=A4)
        product_index = 0
        
        for page_num in range(total_pages):
            for row in range(rows):
                for col in range(cols):
                    if product_index >= len(products):
                        break
                    
                    product = products[product_index]
                    
                    # Position for this label
                    x = margin_x + col * (label_width + gap_x)
                    y = page_height - margin_top - label_height - row * (label_height + gap_y)
                    
                    # 1. Subtle, clean cut-guide rounded border
                    c.setStrokeColorRGB(0.78, 0.81, 0.86)
                    c.setLineWidth(0.5)
                    c.roundRect(x, y, label_width, label_height, 2 * mm, stroke=1, fill=0)
                    
                    # 2. Product Name (Bold, centered, cleanly truncated if long)
                    c.setFillColorRGB(0.06, 0.09, 0.16)
                    c.setFont("Helvetica-Bold", 7.5)
                    max_text_w = label_width - 4 * mm
                    display_name = product.name or "Unnamed Product"
                    while display_name and c.stringWidth(display_name + "...", "Helvetica-Bold", 7.5) > max_text_w:
                        display_name = display_name[:-1]
                    if display_name != (product.name or ""):
                        display_name = display_name.strip() + "..."
                    name_w = c.stringWidth(display_name, "Helvetica-Bold", 7.5)
                    c.drawString(x + (label_width - name_w) / 2, y + label_height - 4.5 * mm, display_name)
                    
                    # 3. Sub-info: Category & UOM
                    c.setFillColorRGB(0.40, 0.45, 0.53)
                    c.setFont("Helvetica", 6.2)
                    sub_parts = []
                    if product.uom:
                        sub_parts.append(f"UOM: {product.uom}")
                    if product.account_title:
                        sub_parts.append(product.account_title)
                    sub_text = " • ".join(sub_parts) if sub_parts else "General"
                    while sub_text and c.stringWidth(sub_text + "...", "Helvetica", 6.2) > max_text_w:
                        sub_text = sub_text[:-1]
                    if sub_parts and sub_text != " • ".join(sub_parts):
                        sub_text = sub_text.strip() + "..."
                    sub_w = c.stringWidth(sub_text, "Helvetica", 6.2)
                    c.drawString(x + (label_width - sub_w) / 2, y + label_height - 7.5 * mm, sub_text)
                    
                    # 4. Barcode Image
                    barcode_file = None
                    if product.barcode:
                        if product.company_id:
                            comp_file = Path(current_app.root_path) / "static" / "product_barcodes" / f"company_{product.company_id}" / f"{product.barcode}.png"
                            if comp_file.exists():
                                barcode_file = comp_file
                        if not barcode_file:
                            root_file = Path(current_app.root_path) / "static" / "product_barcodes" / f"{product.barcode}.png"
                            if root_file.exists():
                                barcode_file = root_file
                    
                    img_w = 56 * mm
                    img_h = 22 * mm
                    img_x = x + (label_width - img_w) / 2
                    img_y = y + 2.8 * mm
                    
                    if barcode_file and barcode_file.exists():
                        try:
                            c.drawImage(str(barcode_file), img_x, img_y, 
                                        width=img_w, height=img_h, 
                                        preserveAspectRatio=True, anchor='c')
                        except Exception:
                            c.setStrokeColorRGB(0.85, 0.85, 0.88)
                            c.rect(img_x, img_y, img_w, img_h)
                            c.setFont("Helvetica", 7)
                            c.setFillColorRGB(0.4, 0.4, 0.4)
                            c.drawCentredString(x + label_width / 2, y + 12 * mm, f"Barcode: {product.barcode}")
                    else:
                        c.setStrokeColorRGB(0.85, 0.85, 0.88)
                        c.rect(img_x, img_y, img_w, img_h)
                        c.setFont("Helvetica", 7)
                        c.setFillColorRGB(0.4, 0.4, 0.4)
                        c.drawCentredString(x + label_width / 2, y + 12 * mm, f"Barcode: {product.barcode or 'N/A'}")
                    
                    product_index += 1
                
                if product_index >= len(products):
                    break
            
            # Professional page footer
            c.setFont("Helvetica", 7.5)
            c.setFillColorRGB(0.45, 0.48, 0.55)
            c.drawCentredString(page_width / 2, 5 * mm, f"iCount Inventory — Page {page_num + 1} of {total_pages}")
            
            # Start new page if not the last page
            if page_num < total_pages - 1:
                c.showPage()
        
        # Save PDF
        c.save()
        output.seek(0)
        
        timestamp = datetime.now(ZoneInfo("Asia/Manila")).strftime('%Y%m%d_%H%M%S')
        filename = f"product_barcodes_{timestamp}.pdf"
        
        log_activity(
            action="Export Barcode PDF",
            details=f"Exported barcode PDF with {len(products)} products across {total_pages} pages"
        )
        
        return Response(
            output.getvalue(),
            mimetype="application/pdf",
            headers={"Content-Disposition": f"attachment;filename={filename}"}
        )
        
    except Exception as e:
        log_activity(
            action="Export Barcode PDF Error",
            details=f"Failed to export barcode PDF: {str(e)}"
        )
        flash(f'Error generating barcode PDF: {str(e)}', 'error')
        return redirect(url_for('misc.transaction_history'))


@misc.route('/download_requisition_history', methods=['POST'])
@login_required
@admin_required
def download_requisition_history():
    """Download requisition history as Excel file"""
    try:
        # Get date range from form
        from_date_str = request.form.get('req_from_date')
        to_date_str = request.form.get('req_to_date')
        
        if not from_date_str or not to_date_str:
            flash('Please select both from and to dates.', 'error')
            return redirect(url_for('misc.misc_dashboard'))
        
        # Parse dates
        from_date = datetime.strptime(from_date_str, '%Y-%m-%d').date()
        to_date = datetime.strptime(to_date_str, '%Y-%m-%d').date()
        
        if from_date > to_date:
            flash('From date cannot be later than to date.', 'error')
            return redirect(url_for('misc.misc_dashboard'))
        
        # Convert dates to datetime for proper comparison
        from_datetime = datetime.combine(from_date, datetime.min.time())
        to_datetime = datetime.combine(to_date, datetime.max.time())
        
        # Get approved requisitions within date range
        # Base filtering on date_created instead of date_approved
        if current_user.role == 'Company-Admin' and current_user.company_id:
            # Company-Admin can only see their own company's requisitions
            requisitions = Requisition.query.filter(
                Requisition.status == 'approved',
                Requisition.company_id == current_user.company_id,
                Requisition.date_created >= from_datetime,
                Requisition.date_created <= to_datetime
            ).order_by(Requisition.date_created.desc()).all()
        else:
            # SYS-Admin and WH-Admin can see all requisitions
            requisitions = Requisition.query.filter(
                Requisition.status == 'approved',
                Requisition.date_created >= from_datetime,
                Requisition.date_created <= to_datetime
            ).order_by(Requisition.date_created.desc()).all()
        
        # Extract requisition IDs and approval dates
        requisition_ids = [req.id for req in requisitions]
        # Use date_created for all requisitions
        approval_dates = {req.id: req.date_created for req in requisitions}
        
        if not requisition_ids:
            flash('No approved requisitions found for the selected date range.', 'warning')
            return redirect(url_for('misc.misc_dashboard'))
        
        # Create Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Requisition History"
        
        # Define styles
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        title_font = Font(bold=True, size=14)
        title_alignment = Alignment(horizontal="center", vertical="center")
        
        # Define borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Write title
        ws.merge_cells('A1:F1')
        title_cell = ws.cell(row=1, column=1, value="ICEBERGS WAREHOUSE REQUISITION HISTORY")
        title_cell.font = title_font
        title_cell.alignment = title_alignment
        
        # Write date range
        ws.merge_cells('A2:F2')
        date_range = f"Period: {from_date.strftime('%b/%d/%Y')} to {to_date.strftime('%b/%d/%Y')}"
        date_cell = ws.cell(row=2, column=1, value=date_range)
        date_cell.font = Font(bold=True, size=12)
        date_cell.alignment = title_alignment
        
        # Write header row
        headers = [
            'trxn date', 'Item Description', 'qty', 'payee', 'branch'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Write requisition data
        current_row = 5
        
        for requisition in requisitions:
            approval_date = approval_dates.get(requisition.id)
            if not approval_date:
                continue
            
            # Write requisition header
            ws.merge_cells(f'A{current_row}:E{current_row}')
            # Use date_created for all requisitions
            req_header = f"Requisition #{requisition.id} - {requisition.requested_by} ({requisition.branch or 'N/A'}) - Created: {approval_date.strftime('%m/%d/%Y %I:%M %p')}"
            header_cell = ws.cell(row=current_row, column=1, value=req_header)
            header_cell.font = Font(bold=True, size=11)
            header_cell.fill = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")
            header_cell.border = thin_border
            current_row += 1
            
            for item in requisition.items:
                # Format item description with item_uom_desc instead of name and uom
                item_description = item.product.item_uom_desc or f"{item.product.name} ({item.uom})"
                
                # Format approval date with time
                trxn_date = approval_date.strftime('%m/%d/%Y %I:%M %p')
                
                # Write row data
                row_data = [
                    trxn_date,
                    item_description,
                    item.quantity,
                    requisition.requested_by,
                    requisition.branch or 'N/A'
                ]
                
                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col, value=value)
                    cell.border = thin_border
                    
                    # Format numbers
                    if col == 3:  # qty column
                        if isinstance(value, (int, float)):
                            cell.number_format = '0.00'
                
                current_row += 1
            
            # Add empty row between requisitions
            current_row += 1
        
        # Auto-adjust column widths
        column_widths = [20, 40, 10, 25, 20]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Prepare response
        timestamp = datetime.now(ZoneInfo("Asia/Manila")).strftime('%Y%m%d_%H%M%S')
        filename = f"requisition_history_{from_date.strftime('%Y%m%d')}_to_{to_date.strftime('%Y%m%d')}_{timestamp}.xlsx"
        
        # Log the export
        log_activity(
            action="Export Requisition History",
            details=f"Exported requisition history from {from_date.strftime('%Y-%m-%d')} to {to_date.strftime('%Y-%m-%d')} with {len(requisition_ids)} requisitions"
        )
        
        return Response(
            output.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment;filename={filename}"}
        )
        
    except Exception as e:
        # Log the error
        log_activity(
            action="Export Requisition History Error",
            details=f"Failed to export requisition history: {str(e)}"
        )
        flash(f'Error exporting requisition history: {str(e)}', 'error')
        return redirect(url_for('misc.misc_dashboard'))


@misc.route('/download_all_transactions', methods=['POST'])
@login_required
@admin_required
def download_all_transactions():
    """Download only TRANS IN transactions with supplier name, item uom desc, qty, exp date, transaction date, and payee"""
    try:
        # Get date range from form
        from_date_str = request.form.get('from_date')
        to_date_str = request.form.get('to_date')
        
        if not from_date_str or not to_date_str:
            flash('Please select both from and to dates.', 'error')
            return redirect(url_for('misc.misc_dashboard'))
        
        # Parse dates
        from_date = datetime.strptime(from_date_str, '%Y-%m-%d').date()
        to_date = datetime.strptime(to_date_str, '%Y-%m-%d').date()
        
        if from_date > to_date:
            flash('From date cannot be later than to date.', 'error')
            return redirect(url_for('misc.misc_dashboard'))
        
        # Convert dates to datetime for proper comparison
        from_datetime = datetime.combine(from_date, datetime.min.time())
        to_datetime = datetime.combine(to_date, datetime.max.time())
        
        # Query only TRANS IN transaction history within date range.
        # Always respect the active store/brand toggle. A selected company must
        # scope the export even for SYS-Admin and WH-Admin users; only Global
        # View (no active company) intentionally exports across all companies.
        history_query = ProductHistory.query.join(Product).filter(
            ProductHistory.trxn_code == 'TRANS IN',
            ProductHistory.trxn_date >= from_datetime,
            ProductHistory.trxn_date <= to_datetime
        )
        if current_user.company_id:
            history_query = history_query.filter(Product.company_id == current_user.company_id)
        history_records = history_query.order_by(ProductHistory.trxn_date.desc()).all()
        
        if not history_records:
            flash('No TRANS IN transactions found for the selected date range.', 'warning')
            return redirect(url_for('misc.misc_dashboard'))
        
        # Create Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "TRANS IN Transactions"
        
        # Define styles
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        title_font = Font(bold=True, size=14)
        title_alignment = Alignment(horizontal="center", vertical="center")
        
        # Define borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Write title
        ws.merge_cells('A1:F1')
        title_cell = ws.cell(row=1, column=1, value="ICEBERGS WAREHOUSE TRANS IN TRANSACTIONS")
        title_cell.font = title_font
        title_cell.alignment = title_alignment
        
        # Write date range
        ws.merge_cells('A2:F2')
        date_range = f"Period: {from_date.strftime('%b/%d/%Y')} to {to_date.strftime('%b/%d/%Y')}"
        date_cell = ws.cell(row=2, column=1, value=date_range)
        date_cell.font = Font(bold=True, size=12)
        date_cell.alignment = title_alignment
        
        # Write header row with required columns
        headers = [
            'Transaction Date', 'Payee', 'Supplier Name', 'Item Description', 'Quantity', 'Expiry Date'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Write transaction data
        current_row = 5
        
        for record in history_records:
            # Get product details
            product = Product.query.get(record.product_id)
            if not product:
                continue
            
            # Get supplier name and expiry date from the referenced batch if available
            supplier_name = 'N/A'
            exp_date = 'N/A'
            
            # If this transaction references a batch, get the batch details
            if record.reference_type == 'batch' and record.reference_id:
                batch = WarehouseProductBatch.query.get(record.reference_id)
                if batch:
                    supplier_name = batch.supplier_name or 'N/A'
                    exp_date = batch.exp_date.strftime('%m/%d/%Y') if batch.exp_date else 'N/A'
            
            # If we couldn't get batch info, try to get from product
            if supplier_name == 'N/A':
                # Get the most recent batch for this product to get supplier info
                recent_batch = WarehouseProductBatch.query.filter_by(product_id=record.product_id).order_by(WarehouseProductBatch.added_on.desc()).first()
                if recent_batch:
                    supplier_name = recent_batch.supplier_name or 'N/A'
                    exp_date = recent_batch.exp_date.strftime('%m/%d/%Y') if recent_batch.exp_date else 'N/A'
            
            # Format item description with item_uom_desc
            item_description = product.item_uom_desc or f"{product.name} ({product.uom or 'unit'})"
            
            # Format transaction date
            trxn_date = record.trxn_date.strftime('%m/%d/%Y %I:%M %p')
            
            # Format payee
            payee = record.payee or 'N/A'
            
            # Format quantity
            quantity = record.quantity_in
            
            # Write row data with required columns
            row_data = [
                trxn_date,
                payee,
                supplier_name,
                item_description,
                quantity,
                exp_date
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col, value=value)
                cell.border = thin_border
                
                # Format numbers
                if col == 5:  # qty column
                    if isinstance(value, (int, float)):
                        cell.number_format = '0.00'
            
            current_row += 1
        
        # Auto-adjust column widths
        column_widths = [20, 25, 25, 40, 15, 15]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Prepare response
        timestamp = datetime.now(ZoneInfo("Asia/Manila")).strftime('%Y%m%d_%H%M%S')
        filename = f"trans_in_transactions_{from_date.strftime('%Y%m%d')}_to_{to_date.strftime('%Y%m%d')}_{timestamp}.xlsx"
        
        # Log the export
        log_activity(
            action="Export TRANS IN Transactions",
            details=f"Exported TRANS IN transactions from {from_date.strftime('%Y-%m-%d')} to {to_date.strftime('%Y-%m-%d')} with {len(history_records)} records"
        )
        
        return Response(
            output.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment;filename={filename}"}
        )
        
    except Exception as e:
        # Log the error
        log_activity(
            action="Export TRANS IN Transactions Error",
            details=f"Failed to export TRANS IN transactions: {str(e)}"
        )
        flash(f'Error exporting TRANS IN transactions: {str(e)}', 'error')
        return redirect(url_for('misc.misc_dashboard'))


@misc.route('/download_batch_barcodes')
@login_required
@admin_required
def download_batch_barcodes():
    """Download batch barcodes has been retired."""
    flash('Batch barcodes have been retired. Please use product barcodes.', 'info')
    return redirect(url_for('misc.misc_dashboard'))


@misc.route('/print_batch_barcodes', methods=['GET', 'POST'])
@login_required
@admin_required
def print_batch_barcodes():
    """Print batch barcodes has been retired."""
    flash('Batch barcodes have been retired. Please use product barcodes.', 'info')
    return redirect(url_for('misc.misc_dashboard'))
