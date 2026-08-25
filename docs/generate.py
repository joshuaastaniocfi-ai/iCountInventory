import os
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from datetime import datetime, timedelta

# ====================
# CONFIG
# ====================
OUTPUT_DIR = "output"
os.makedirs(OUTPUT_DIR, exist_ok=True)

PREPARED_BY = "Your Name"
START_DATE = datetime(2025, 9, 8)
RATE_PER_DAY = 800

PHASES = [
    {"name": "Phase 1: Requirements & DB Schema", "duration": 15},
    {"name": "Phase 2: Product Catalog & Barcode", "duration": 20},
    {"name": "Phase 3: Ordering & Shipment", "duration": 25},
    {"name": "Phase 4: Expiration Pre-Alerts", "duration": 15},
    {"name": "Phase 5: Sales Integration", "duration": 20},
    {"name": "Phase 6: Goods Receipt", "duration": 15},
    {"name": "Phase 7: Inventory Counting", "duration": 15},
    {"name": "Phase 8: Item Creation & Control", "duration": 10},
    {"name": "Phase 9: Discounts & Promos", "duration": 10},
    {"name": "Phase 10: Spoilage & Transfers", "duration": 10},
    {"name": "Phase 11: Data Adjustments", "duration": 5},
    {"name": "Phase 12: Exporting & Variance Reports", "duration": 10}
]

# ====================
# Gantt Chart (Static Image)
# ====================
def create_gantt(phases, out_path):
    start = START_DATE
    fig, ax = plt.subplots(figsize=(10, 6))
    ylabels, starts, ends = [], [], []

    for i, p in enumerate(phases):
        phase_start = start
        phase_end = start + timedelta(days=p["duration"])
        ylabels.append(p["name"])
        starts.append(phase_start)
        ends.append(phase_end)
        ax.barh(i, (phase_end - phase_start).days, left=phase_start, color="skyblue")
        start = phase_end

    ax.set_yticks(range(len(ylabels)))
    ax.set_yticklabels(ylabels, fontsize=8)
    ax.invert_yaxis()
    ax.set_xlabel("Timeline")
    plt.tight_layout()
    plt.savefig(out_path)
    plt.close()

# ====================
# Word Export with Editable Table + Static Chart
# ====================
def export_word(phases, gantt_img, out_path):
    doc = Document()

    # Title
    doc.add_heading("Retail Operations Workflow – iCount System", 0)
    doc.add_paragraph(f"Prepared by: {PREPARED_BY}\nDate: {datetime.today().strftime('%B %d, %Y')}")

    # Executive Summary
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("This proposal outlines the development of the iCount system...")

    # Editable Gantt Table
    doc.add_heading("Project Timeline (Editable Gantt Table)", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid"

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Phase"
    hdr_cells[1].text = "Start"
    hdr_cells[2].text = "End"
    hdr_cells[3].text = "Duration (days)"
    hdr_cells[4].text = "Timeline Bar"

    start = START_DATE
    for p in phases:
        row_cells = table.add_row().cells
        end = start + timedelta(days=p["duration"])
        row_cells[0].text = p["name"]
        row_cells[1].text = start.strftime("%b %d, %Y")
        row_cells[2].text = end.strftime("%b %d, %Y")
        row_cells[3].text = str(p["duration"])
        # Create a simple text bar
        row_cells[4].text = "█" * (p["duration"] // 2)
        start = end

    doc.add_paragraph("Note: You can adjust the 'Timeline Bar' column manually by resizing columns or editing blocks.")

    # Static Gantt Image
    doc.add_heading("Project Timeline (Static Chart)", level=1)
    doc.add_picture(gantt_img, width=Inches(6))

    # Costing Pie Chart
    doc.add_heading("Project Cost Breakdown", level=1)
    doc.add_picture(os.path.join(OUTPUT_DIR, "icount_cost_pie.png"), width=Inches(4))

    # Save document
    doc.save(out_path)

# ====================
# Cost Pie Chart
# ====================
def create_pie(phases, out_path):
    sizes = [p["duration"] for p in phases]
    labels = [p["name"] for p in phases]

    fig, ax = plt.subplots(figsize=(6, 6))
    ax.pie(sizes, labels=labels, autopct="%1.1f%%", startangle=140)
    ax.set_title("Cost Distribution by Phase")
    plt.tight_layout()
    plt.savefig(out_path)
    plt.close()

# ====================
# Main
# ====================
def main():
    gantt_path = os.path.join(OUTPUT_DIR, "icount_gantt.png")
    pie_path = os.path.join(OUTPUT_DIR, "icount_cost_pie.png")
    word_out = os.path.join(OUTPUT_DIR, "iCount_Proposal.docx")

    print("Generating Gantt chart...")
    create_gantt(PHASES, gantt_path)

    print("Generating Pie chart...")
    create_pie(PHASES, pie_path)

    print("Exporting Word with editable table + static chart...")
    export_word(PHASES, gantt_path, word_out)

    print("DONE. File saved at:", word_out)

if __name__ == "__main__":
    main()
